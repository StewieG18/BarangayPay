from docx import Document
from docx.shared import Pt

TEMPLATE = r"C:\Users\Jhouanese\Downloads\gcash-frd-specification.docx"
OUTPUT = r"C:\Users\Jhouanese\Desktop\BarangayPay\BarangayPay-FRD.docx"


def set_cell(cell, value):
    cell.text = str(value)
    for p in cell.paragraphs:
        for run in p.runs:
            run.font.name = "Arial"
            run.font.size = Pt(9)


def set_table(table, headers, rows):
    while len(table.rows) > 1:
        table._tbl.remove(table.rows[-1]._tr)
    for i, value in enumerate(headers):
        set_cell(table.rows[0].cells[i], value)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell(cells[i], value)


doc = Document(TEMPLATE)
doc.core_properties.title = "BarangayPay Functional Requirement Document"
doc.core_properties.subject = "Functional requirements for the BarangayPay resident service portal"
doc.core_properties.author = "BarangayPay Product and QA Team"

doc.paragraphs[0].text = "FUNCTIONAL REQUIREMENT DOCUMENT (FRD)\nBarangayPay Resident Services and Payment Portal"
doc.paragraphs[2].text = "1. Persona and User Story Mapping"
doc.paragraphs[4].text = "2. Step by Step Functional Workflow Execution"
doc.paragraphs[6].text = "3. Business Rules and Validation Logic Matrix"
doc.paragraphs[8].text = "4. Exception and Error Handling Matrix"

set_table(doc.tables[0], ["Document ID", "FRD-BP-RESIDENT-2026-V1.0"], [
    ["Parent BRD Reference", "BRD-BP-RESIDENT-2026-V1.0"],
    ["Feature Module", "Resident login, barangay service requests, document submission, status tracking, and applicable payments"],
    ["Primary Author", "BarangayPay Product and QA Team"],
    ["Source System", "BarangayPay web application resident portal"],
])

set_table(doc.tables[1], ["User Story ID", "User Persona", "User Need / Action", "Expected Business Benefit"], [
    ["US-BP-101", "Barangay Resident", "Sign in securely with an email address and password to access the resident portal.", "Protects account information and provides a personalized service workspace."],
    ["US-BP-102", "Barangay Resident", "Browse available services and submit a request for a clearance or certificate online.", "Reduces in-person visits and makes barangay applications easier to start."],
    ["US-BP-103", "Barangay Resident", "Enter contact details, purpose, preferred date, notes, and supporting documents for a selected service.", "Improves completeness of applications and gives barangay staff the information needed to process them."],
    ["US-BP-104", "Barangay Resident", "Review a request, confirm its accuracy, monitor its status, and pay an applicable fee through the payment provider.", "Provides transparency from submission through completion and supports convenient digital payment."],
])

set_table(doc.tables[2], ["Step #", "App Screen / View", "User Action / Input", "System Processing and State Change"], [
    ["Step 1", "Login", "Enter email address and password, then select Sign In.", "Validate required credentials and authenticate the resident; show an error when authentication fails."],
    ["Step 2", "Dashboard", "View transaction summary and select Request a Service, Services, or My Transactions.", "Load the resident dashboard and display pending, processing, completed, and payment summaries."],
    ["Step 3", "Available Services", "Choose Barangay Clearance, Certificate of Residency, Certificate of Indigency, or Barangay Business Clearance.", "Open the service request workflow with the selected service preselected."],
    ["Step 4", "Personal Information", "Enter first name, last name, optional middle name, contact number, and address.", "Validate required fields and retain the information as part of the draft request."],
    ["Step 5", "Request Details", "Select purpose, choose a preferred date, and optionally enter request notes.", "Validate the purpose and date, then associate the details with the selected service."],
    ["Step 6", "Supporting Documents", "Upload one or more PDF, JPG, JPEG, or PNG files.", "Validate file type and configured size limits, store the files, and associate them with the draft request."],
    ["Step 7", "Review Your Request", "Review the information and select the confirmation checkbox.", "Prevent submission until the resident confirms that the information is accurate and complete."],
    ["Step 8", "Submission, Payment, and Tracking", "Submit the request and, when a fee applies, complete payment through the configured payment provider.", "Create a transaction reference, set the request to Pending, record payment status, and show the request in My Transactions."],
])

set_table(doc.tables[3], ["Rule ID", "Business Rule Name", "Validation Logic and Condition", "System Action on Violation"], [
    ["BR-BP-01", "Credential Requirement", "Email address and password are required for sign in.", "Keep the resident on the login page and identify the missing or invalid credential."],
    ["BR-BP-02", "Service Selection", "A request must contain one supported service from the configured service catalog.", "Block the workflow until a valid service is selected."],
    ["BR-BP-03", "Required Resident Information", "First name, last name, contact number, and address are required; middle name is optional.", "Display field-level validation and prevent progression until required values are complete."],
    ["BR-BP-04", "Purpose and Date", "Purpose must be selected from the configured list and preferred date must be a valid selectable date.", "Reject invalid values and ask the resident to correct the request details."],
    ["BR-BP-05", "Document Validation", "Uploaded requirements must use PDF, JPG, JPEG, or PNG format and remain within the configured file-size limit.", "Reject unsupported or oversized files and keep valid uploads attached to the request."],
    ["BR-BP-06", "Submission Confirmation", "The resident must confirm that the information is accurate and complete before submission.", "Disable or reject submission until the confirmation checkbox is selected."],
    ["BR-BP-07", "Payment Applicability", "Only services configured with a fee require payment; fee-free services proceed without checkout.", "Create the correct payment state and show the applicable next action."],
    ["BR-BP-08", "Request Status", "A submitted request starts as Pending and may move through Processing to Completed or Rejected.", "Record each status transition with its transaction reference and make it visible in My Transactions."],
])

set_table(doc.tables[4], ["Error Code", "Trigger Scenario", "Backend System Behavior", "Client UI Display Message"], [
    ["ERR-BP-101", "Required login field is empty", "Reject the request before authentication.", "Please enter your email address and password."],
    ["ERR-BP-102", "Email or password is invalid", "Return an authentication failure without revealing which credential was incorrect.", "The email or password is incorrect. Please try again."],
    ["ERR-BP-103", "Required request information is missing", "Reject draft submission and return field validation details.", "Complete all required information before continuing."],
    ["ERR-BP-104", "Unsupported or oversized document upload", "Do not attach the invalid file and retain the remaining valid files.", "This file type or size is not accepted. Upload a PDF, JPG, JPEG, or PNG within the allowed limit."],
    ["ERR-BP-105", "Confirmation checkbox is not selected", "Reject submission without creating a transaction.", "Please confirm that your information is accurate and complete."],
    ["ERR-BP-106", "Payment provider returns failure or timeout", "Mark payment as Failed or Pending and keep the request reference available for status checking.", "Payment could not be completed. Check My Transactions before trying again."],
    ["ERR-BP-107", "Service request submission fails", "Do not create a duplicate request; log the failure and return a retry-safe response.", "We could not submit your request. Please try again."],
])

doc.save(OUTPUT)
print(OUTPUT)
