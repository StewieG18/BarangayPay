from docx import Document
from docx.shared import Pt

src = r"C:\Users\Jhouanese\Downloads\DATABASE NORMALIZATION_PRECILLASDB.docx"
dst = r"C:\Users\Jhouanese\Documents\From Data Form to Normalized Database(Condor).docx"

doc = Document(src)

# Add fields visible on the supplied enrollment form to the field list.
doc.paragraphs[4].insert_paragraph_before("Form No.")
doc.paragraphs[5].insert_paragraph_before("Academic Year")
doc.paragraphs[18].insert_paragraph_before("Student Signature")
doc.paragraphs[19].insert_paragraph_before("Registrar Signature")

# Remove the stray character before the 3NF heading.
for p in doc.paragraphs:
    if p.text.strip() == "t":
        p.text = ""

# Make the final junction-table heading agree with the actual columns and ERD.
for p in doc.paragraphs:
    if "ENROLLED_SUBJECTS" in p.text and "Final Normalized" not in p.text:
        if "Subject_Code" in p.text:
            p.text = "ENROLLED_SUBJECTS - PK: (Enrollment_ID, Offering_ID) | FK: Enrollment_ID, Offering_ID"

# Correct the explanatory FK list to refer to offerings in the final design.
for p in doc.paragraphs:
    if p.text.startswith("Enforced referential integrity:"):
        p.text = p.text.replace("Enrollment_ID, Subject_Code", "Enrollment_ID, Offering_ID")

# Correct the sample offering identifier in the final junction table.
if len(doc.tables) > 13:
    rows = doc.tables[13].rows
    if len(rows) > 2:
        rows[2].cells[2].text = "OFF-0001"

doc.save(dst)
print(dst)
